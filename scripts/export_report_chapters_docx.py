"""Export report chapters from Markdown to DOCX.

This script is intentionally scoped to the thesis report Markdown files in
``reports/``. It uses python-docx from the project virtual environment.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
TEMPLATE = ROOT / "templates" / "ĐATN_NguyenQuangSang - V1.docx"
OUTPUT = REPORTS / "bao-cao-chuong-3-4-5.docx"
IMAGE_CACHE = REPORTS / "assets" / "docx"

CHAPTERS = [
    REPORTS / "chuong-3-phan-tich-thiet-ke-he-thong.md",
    REPORTS / "chuong-4-xay-dung-he-thong.md",
    REPORTS / "chuong-5-thuc-nghiem-va-danh-gia.md",
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text.strip())
    run.bold = bold
    for run_ in paragraph.runs:
        run_.font.name = "Times New Roman"
        run_.font.size = Pt(10)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline_markdown(paragraph, text: str) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
        elif token.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])
    for run in paragraph.runs:
        if run.font.name is None:
            run.font.name = "Times New Roman"
        if run.font.size is None:
            run.font.size = Pt(12)


def resolve_image(path_text: str) -> Path:
    path = Path(path_text)
    if not path.is_absolute():
        path = REPORTS / path
    return path


def crop_png(png_path: Path) -> Path:
    IMAGE_CACHE.mkdir(parents=True, exist_ok=True)
    out = IMAGE_CACHE / f"{png_path.stem}.cropped.png"
    if out.exists() and out.stat().st_mtime >= png_path.stat().st_mtime:
        return out

    image = Image.open(png_path).convert("RGB")
    pixels = image.load()
    width, height = image.size
    min_x, min_y = width, height
    max_x, max_y = 0, 0
    for y in range(height):
        for x in range(width):
            r, g, b = pixels[x, y]
            if min(r, g, b) < 248:
                min_x = min(min_x, x)
                min_y = min(min_y, y)
                max_x = max(max_x, x)
                max_y = max(max_y, y)
    if min_x >= max_x or min_y >= max_y:
        image.save(out)
        return out
    pad = 24
    box = (
        max(min_x - pad, 0),
        max(min_y - pad, 0),
        min(max_x + pad, width),
        min(max_y + pad, height),
    )
    image.crop(box).save(out)
    return out


def image_for_docx(image_path: Path) -> Path:
    if image_path.suffix.lower() == ".svg":
        rel = image_path.relative_to(REPORTS)
        ql_png = REPORTS / str(rel).replace("assets/mermaid/", "assets/mermaid-png/")
        ql_png = Path(f"{ql_png}.png")
        if not ql_png.exists():
            raise FileNotFoundError(
                f"Missing PNG preview for {image_path}. Run qlmanage conversion first."
            )
        return crop_png(ql_png)
    return image_path


def add_picture(doc: Document, image_path: Path, alt_text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    usable_width = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    width = min(usable_width, Inches(6.2))
    run.add_picture(str(image_for_docx(image_path)), width=width)
    if alt_text:
        paragraph.paragraph_format.space_after = Pt(3)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for r_idx, row in enumerate(rows):
        for c_idx in range(cols):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            set_cell_text(cell, text, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shading(cell, "D9EAF7")
    doc.add_paragraph()


def add_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    add_inline_markdown(paragraph, text)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text.strip("*"))
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    paragraph.paragraph_format.space_after = Pt(2)


def add_source(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline_markdown(paragraph, text)
    for run in paragraph.runs:
        run.italic = True
        run.font.size = Pt(10)
    paragraph.paragraph_format.space_after = Pt(8)


def add_report_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8 if level == 1 else 6)
    paragraph.paragraph_format.space_after = Pt(6 if level <= 2 else 4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt({1: 16, 2: 14}.get(level, 13))


def consume_equation(lines: list[str], start: int) -> tuple[str, int]:
    buf = [lines[start]]
    i = start + 1
    while i < len(lines):
        buf.append(lines[i])
        if lines[i].strip() == "$$":
            return "\n".join(buf), i + 1
        i += 1
    return "\n".join(buf), i


def export_markdown_file(doc: Document, path: Path, first: bool) -> None:
    if not first:
        doc.add_page_break()

    lines = path.read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "$$":
            equation, i = consume_equation(lines, i)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(equation)
            run.font.name = "Cambria Math"
            run.font.size = Pt(11)
            continue

        image_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            alt, image_path = image_match.groups()
            add_picture(doc, resolve_image(image_path), alt)
            i += 1
            continue

        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            add_report_heading(doc, text, min(level, 3))
            i += 1
            continue

        if stripped.startswith("**Hình ") or stripped.startswith("**Bảng "):
            add_caption(doc, stripped)
            i += 1
            continue

        if stripped.startswith("Nguồn:"):
            add_source(doc, stripped)
            i += 1
            continue

        if stripped.startswith("- "):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            paragraph.paragraph_format.first_line_indent = Inches(-0.15)
            add_inline_markdown(paragraph, f"• {stripped[2:]}")
            i += 1
            continue

        add_paragraph(doc, stripped)
        i += 1


def configure_base_styles(doc: Document) -> None:
    section = doc.sections[-1]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(1)

    styles = doc.styles
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3"]:
        style = styles[name]
        style.font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(12)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 2"].font.bold = True
    styles["Heading 3"].font.size = Pt(13)
    styles["Heading 3"].font.bold = True


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def main() -> None:
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    configure_base_styles(doc)

    for idx, chapter in enumerate(CHAPTERS):
        export_markdown_file(doc, chapter, first=(idx == 0))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
