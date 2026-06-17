"""DOCX rendering for generated administrative documents."""

import re
from io import BytesIO

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


BODY_FONT_SIZE = 13
TITLE_MARKERS = ("CÔNG VĂN", "TỜ TRÌNH", "QUYẾT ĐỊNH", "THÔNG BÁO", "KẾ HOẠCH", "BIÊN BẢN")


def render_administrative_docx(content: str) -> bytes:
    """Render a drafted administrative document into DOCX bytes."""
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Times New Roman"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal_style.font.size = Pt(BODY_FONT_SIZE)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(0)

    lines = _render_header_table(document, content.splitlines())
    lines = _render_document_metadata(document, lines)
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)

        if _looks_like_title(line):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line)
            run.bold = True
            run.font.size = Pt(14)
        elif _looks_like_subtitle(line):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line)
            run.bold = True
        elif _looks_like_signature_block(line):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(line)
            run.bold = True
        elif _looks_like_place_and_date(line):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = paragraph.add_run(line)
            run.italic = True
        elif _looks_like_section_heading(line):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1)
            run = paragraph.add_run(line)
            run.bold = True
        elif line.lower().startswith("nơi nhận"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line)
            run.bold = True
            run.italic = True
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if not _looks_like_list_item(line) and not line.lower().startswith(("số:", "kính gửi:")):
                paragraph.paragraph_format.first_line_indent = Cm(1)
            paragraph.add_run(line)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _render_header_table(document: Document, lines: list[str]) -> list[str]:
    """Render agency and national headers as the standard two-column block."""
    national_index = next(
        (index for index, line in enumerate(lines[:10]) if _is_national_header(line)),
        None,
    )
    if national_index is None or national_index == 0:
        return lines

    left_lines = [line.strip() for line in lines[:national_index] if line.strip()]
    if not left_lines:
        return lines

    right_end = national_index + 1
    while right_end < len(lines) and right_end <= national_index + 3:
        candidate = lines[right_end].strip()
        if not candidate:
            break
        if _looks_like_document_number(candidate) or _looks_like_place_and_date(candidate):
            break
        right_end += 1
    right_lines = [line.strip() for line in lines[national_index:right_end] if line.strip()]

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(7.5)
    table.columns[1].width = Cm(8.5)
    _remove_table_borders(table)

    for cell, header_lines in zip(table.rows[0].cells, (left_lines, right_lines), strict=True):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell.text = ""
        for index, header_line in enumerate(header_lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(header_line)
            run.bold = True
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(12.5)

    remaining = lines[right_end:]
    while remaining and not remaining[0].strip():
        remaining.pop(0)
    return remaining


def _remove_table_borders(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def _render_document_metadata(document: Document, lines: list[str]) -> list[str]:
    """Place document number and issue date on the same two-column row."""
    significant = [(index, line.strip()) for index, line in enumerate(lines[:8]) if line.strip()]
    number_item = next(
        ((index, line) for index, line in significant if _looks_like_document_number(line)),
        None,
    )
    date_item = next(
        ((index, line) for index, line in significant if _looks_like_place_and_date(line)),
        None,
    )
    if number_item is None or date_item is None:
        return lines

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(7.5)
    table.columns[1].width = Cm(8.5)
    _remove_table_borders(table)

    for cell, text, alignment in (
        (table.cell(0, 0), number_item[1], WD_ALIGN_PARAGRAPH.LEFT),
        (table.cell(0, 1), date_item[1], WD_ALIGN_PARAGRAPH.RIGHT),
    ):
        paragraph = cell.paragraphs[0]
        paragraph.alignment = alignment
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(text)
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(BODY_FONT_SIZE)
        if alignment == WD_ALIGN_PARAGRAPH.RIGHT:
            run.italic = True

    removed_indices = {number_item[0], date_item[0]}
    remaining = [line for index, line in enumerate(lines) if index not in removed_indices]
    while remaining and not remaining[0].strip():
        remaining.pop(0)
    return remaining


def _is_national_header(line: str) -> bool:
    upper = line.upper()
    return (
        "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in upper
        or "CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM" in upper
    )


def _looks_like_title(line: str) -> bool:
    upper = line.upper()
    return line == upper and any(marker in upper for marker in TITLE_MARKERS)


def _looks_like_subtitle(line: str) -> bool:
    return line.lower().startswith(("về việc", "v/v"))


def _looks_like_signature_block(line: str) -> bool:
    upper = line.upper()
    return upper.startswith(
        ("THỦ TRƯỞNG", "HIỆU TRƯỞNG", "GIÁM ĐỐC", "TRƯỞNG PHÒNG", "NGƯỜI KÝ")
    )


def _looks_like_place_and_date(line: str) -> bool:
    lower = line.lower()
    return "ngày" in lower and "tháng" in lower and "năm" in lower and len(line) < 140


def _looks_like_document_number(line: str) -> bool:
    return line.lower().startswith("số:")


def _looks_like_section_heading(line: str) -> bool:
    return bool(re.match(r"^\d+[.)]\s+\S", line))


def _looks_like_list_item(line: str) -> bool:
    return bool(re.match(r"^(?:[-–—•]|[a-zđ][.)]|\(\w+\))\s+", line, flags=re.IGNORECASE))


__all__ = ["render_administrative_docx"]
