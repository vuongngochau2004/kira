"""Extractor for PowerPoint presentations (.pptx)."""

import logging

from src.modules.document.domain.models import ExtractionResult

logger = logging.getLogger(__name__)


def extract_pptx(file_path: str) -> ExtractionResult:
    """Extract text from PPTX using python-pptx.

    Args:
        file_path: Path to PPTX file

    Returns:
        ExtractionResult with text and metadata
    """
    logger.info("Extracting PPTX using python-pptx: %s", file_path)
    try:
        from pptx import Presentation

        prs = Presentation(file_path)
        text_parts = []

        for slide in prs.slides:
            slide_text = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text.append(shape.text)
            if slide_text:
                text_parts.append("\n".join(slide_text))

        full_text = "\n\n".join(text_parts)

        return ExtractionResult(
            text=full_text,
            pages=len(prs.slides),
            metadata={"extractor": "python-pptx", "extraction_method": "normal", "ocr_used": False},
            success=bool(full_text.strip()),
        )
    except Exception as e:
        logger.error("PPTX extraction failed for %s: %s", file_path, e)
        return ExtractionResult(
            success=False,
            error=f"PPTX extraction failed: {e}",
        )


__all__ = ["extract_pptx"]
