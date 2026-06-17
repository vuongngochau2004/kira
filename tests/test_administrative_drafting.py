from io import BytesIO

from docx import Document

from src.modules.chat.infrastructure.handlers.drafting_handler import AdministrativeDraftingHandler
from src.modules.drafting.application.docx_renderer import render_administrative_docx
from src.modules.drafting.application.postprocess import clean_administrative_draft
from src.modules.drafting.domain.prompts.administrative import build_administrative_drafting_prompt


def test_prompt_requires_coherent_draft_without_internal_citations() -> None:
    prompt = build_administrative_drafting_prompt("Soạn thông báo", "[Document 1] Nội dung")

    assert "không phải bản tóm tắt tài liệu" in prompt
    assert "Tuyệt đối không đưa mã nguồn nội bộ" in prompt
    assert 'Không thêm mục "Tài liệu/căn cứ đã sử dụng"' in prompt


def test_clean_draft_removes_document_markers_and_trailing_source_list() -> None:
    content = """THÔNG BÁO
Về việc thanh toán giờ giảng

Nội dung áp dụng [Document 2].

Tài liệu/căn cứ đã sử dụng: [Document 2], [Document 4].
"""

    cleaned = clean_administrative_draft(content)

    assert "[Document" not in cleaned
    assert "Tài liệu/căn cứ đã sử dụng" not in cleaned
    assert cleaned.endswith("Nội dung áp dụng.")


def test_parser_cleans_citations_before_returning_draft() -> None:
    response = """{
      "can_draft": true,
      "reason": "Đủ căn cứ",
      "missing_info": [],
      "draft": "THÔNG BÁO\\nNội dung [Document 1]."
    }"""

    parsed = AdministrativeDraftingHandler._parse_drafting_response(response)

    assert parsed["can_draft"] is True
    assert parsed["draft"] == "THÔNG BÁO\nNội dung."


def test_docx_renderer_creates_two_column_administrative_header() -> None:
    content = """ĐẠI HỌC ĐÀ NẴNG
TRƯỜNG ĐẠI HỌC BÁCH KHOA

CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM
Độc lập - Tự do - Hạnh phúc
-------------------

Số: [CẦN BỔ SUNG]
Đà Nẵng, ngày [CẦN BỔ SUNG], tháng [CẦN BỔ SUNG], năm [CẦN BỔ SUNG]

THÔNG BÁO
Về việc thanh toán giờ giảng

1. Đối tượng áp dụng
Nội dung thực hiện.

HIỆU TRƯỞNG
"""

    rendered = Document(BytesIO(render_administrative_docx(content)))

    assert len(rendered.tables) == 2
    assert "TRƯỜNG ĐẠI HỌC BÁCH KHOA" in rendered.tables[0].cell(0, 0).text
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in rendered.tables[0].cell(0, 1).text
    assert "Số:" in rendered.tables[1].cell(0, 0).text
    assert "Đà Nẵng" in rendered.tables[1].cell(0, 1).text
    assert any(paragraph.text == "THÔNG BÁO" for paragraph in rendered.paragraphs)
